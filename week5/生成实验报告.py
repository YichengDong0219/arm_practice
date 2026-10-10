"""生成实验三和实验四 Word 实验报告。"""

from __future__ import annotations

import json
import sys
from pathlib import Path

TOOLS = Path(__file__).resolve().parents[1] / ".report_tools"
sys.path.insert(0, str(TOOLS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "实验报告素材"
OUTPUT = ROOT / "实验三与实验四实验报告.docx"

NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F8FC"
GRAY = "667085"
LIGHT_GRAY = "D9D9D9"
BLACK = "000000"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=11, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None, numeric_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    borders = table_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), LIGHT_GRAY)

    numeric_cols = set(numeric_cols or [])
    for row_index, row in enumerate(table.rows):
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, east_asia="微软雅黑", latin="Arial", size=9.5, bold=True)
                        run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        set_run_font(run, size=9.5)

    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths, numeric_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = str(header)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    style_table(table, widths=widths, numeric_cols=numeric_cols)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text, first_line=True, bold_lead=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    if first_line:
        p.paragraph_format.first_line_indent = Pt(22)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    return p


def set_picture_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(doc, filename, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    set_picture_alt_text(shape, caption, caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run(caption)
    set_run_font(r, size=9.5)
    return cap


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.83)
    section.right_margin = Inches(0.83)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_specs = {
        "Heading 1": (16, "微软雅黑", 12, 6),
        "Heading 2": (13, "微软雅黑", 9, 4),
        "Heading 3": (11.5, "微软雅黑", 7, 3),
    }
    for name, (size, font, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("五自由度舵机机械臂运动学建模与正方形轨迹仿真实验报告")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=25, bold=True)
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(36)
    r = subtitle.add_run("实验三与实验四")
    set_run_font(r, east_asia="微软雅黑", latin="Arial", size=16)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.65)
    summary.paragraph_format.right_indent = Inches(0.65)
    summary.paragraph_format.line_spacing = 1.5
    r = summary.add_run(
        "基于实测标准 D-H 参数完成五自由度机械臂建模，"
        "并利用五次时间缩放和带边界数值逆运动学实现末端水平正方形轨迹仿真"
    )
    set_run_font(r, size=12)

    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Python 3.10  ·  Robotics Toolbox 1.4.4  ·  SciPy 1.15.3")
    set_run_font(r, size=10)
    doc.add_page_break()


def add_code_block(doc, lines):
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_cell = OxmlElement("w:shd")
        set_cell.set(qn("w:fill"), "F2F4F7")
        p._p.get_or_add_pPr().append(set_cell)
        r = p.add_run(line if line else " ")
        set_run_font(r, east_asia="等线", latin="Consolas", size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    metrics = json.loads((ASSETS / "实验四数值结果.json").read_text(encoding="utf-8"))
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("摘要", level=1)
    add_body(
        doc,
        "本报告完成五自由度舵机机械臂的实测标准 D-H 建模与末端正方形轨迹仿真。实验三根据实体关节轴线和连杆中心距整理五组 D-H 参数，构建 DHRobot 模型并检查零输入姿态。实验四复用同一组参数，以初始末端位置为起点生成边长 35 mm 的水平正方形，四条边各取 70 个采样点，并采用五次时间缩放实现边段平滑启停。逐点逆运动学由带边界最小二乘法完成，只约束末端位置，同时加入相邻关节角连续性项。实际运行得到 280 个有效轨迹点，最大位置误差为 7.9565×10⁻⁷ mm，闭合误差为 7.9564×10⁻⁷ mm，相邻采样点最大关节变化为 0.536826°。结果表明目标路径与仿真末端路径高度重合，所有逆解均满足关节限位，轨迹连续且闭合。",
    )
    p = doc.add_paragraph()
    r = p.add_run("关键词：")
    set_run_font(r, bold=True)
    r = p.add_run("标准 D-H 参数；正运动学；笛卡尔轨迹规划；五次时间缩放；数值逆运动学")
    set_run_font(r)

    doc.add_heading("1 实验概述", level=1)
    doc.add_heading("1.1 实验目标", level=2)
    goals = [
        "根据实体五自由度串联舵机机械臂的转轴拓扑和实测尺寸建立标准 D-H 模型。",
        "验证实测尺寸、关节零位补偿与虚拟机械臂骨架之间的一致性。",
        "在笛卡尔空间规划固定高度的水平正方形闭合路径，并掌握五次时间缩放。",
        "使用带边界最小二乘法逐点求解只约束末端位置的逆运动学。",
        "通过关节限位、位置误差、关节连续性和闭合误差评价轨迹质量。",
    ]
    for goal in goals:
        add_bullet(doc, goal)

    doc.add_heading("1.2 实验环境", level=2)
    add_table(
        doc,
        ["类别", "配置"],
        [
            ["硬件对象", "五自由度串联舵机机械臂；本报告结果为软件仿真"],
            ["软件环境", "Python 3.10.21，Conda 环境 ra_class"],
            ["核心库", "NumPy 2.2.6，SciPy 1.15.3，Matplotlib 3.10.9，Robotics Toolbox 1.4.4"],
            ["关联代码", "实验三.py、实验四.py、生成实验四报告素材.py"],
            ["长度与角度单位", "D-H 长度使用 mm；字典中的角度使用 deg，传入工具箱时转换为 rad"],
        ],
        [1.35, 5.25],
    )

    doc.add_heading("2 实验三 舵机机械臂实机 D-H 参数测量与虚拟搭建", level=1)
    doc.add_heading("2.1 建模方法", level=2)
    add_body(
        doc,
        "实验三首先确定五个舵机的物理旋转轴线，再按照标准 D-H 规则确定相邻轴之间的公垂线。底座到第一转轴的竖直高度记为偏距 d，平行俯仰轴之间的中心距记为连杆长度 a，腕部轴线与末端夹持中心之间的距离记为末端偏距。扭角 α 用于描述相邻转轴方向的变化，offset 用于把工具箱中的数学零位映射到机械臂的工程零位。",
    )
    add_body(
        doc,
        "程序将每组参数组织为包含 theta、alpha、d、a 和 offset 的字典，并按关节顺序建立五个 RevoluteDH 连杆。由于现有参数表中的 90 表示 90°，而 Robotics Toolbox 接口采用弧度，仿真分析在构造模型时执行 deg2rad 转换；这一转换只改变接口表示，不改变实测参数、扭角方向或零位补偿定义。",
    )

    doc.add_heading("2.2 实测标准 D-H 参数", level=2)
    add_table(
        doc,
        ["关节", "θ deg", "α deg", "d mm", "a mm", "offset deg"],
        [
            ["J1", "0", "90", "108.50", "0.00", "0"],
            ["J2", "0", "0", "0.00", "105.00", "90"],
            ["J3", "0", "0", "0.00", "97.50", "0"],
            ["J4", "0", "90", "0.00", "28.22", "90"],
            ["J5", "0", "0", "172.83", "0.00", "0"],
        ],
        [0.65, 0.85, 0.85, 1.05, 1.05, 1.15],
        numeric_cols=[1, 2, 3, 4, 5],
    )
    add_body(
        doc,
        "参数的几何含义如下：d₁=108.50 mm 表示底座安装面到第一关节参考面的高度；a₂=105.00 mm 和 a₃=97.50 mm 对应主要臂段中心距；a₄=28.22 mm 描述腕部局部横向偏移；d₅=172.83 mm 描述腕部末端轴到工具中心的轴向距离。J2 和 J4 的 90° offset 用于匹配实机零位。",
    )

    doc.add_heading("2.3 程序实现", level=2)
    add_code_block(
        doc,
        [
            "links = [rtb.RevoluteDH(",
            "    d=p['d'], a=p['a'],",
            "    alpha=np.deg2rad(p['alpha']),",
            "    offset=np.deg2rad(p.get('offset', 0.0)))",
            "    for p in dh_params_servo]",
            "robot = rtb.DHRobot(links, name='Physical_Servo_Manipulator')",
        ],
    )
    add_body(
        doc,
        "模型建立后使用正运动学逐级计算 Frame 0 至 Frame 5 的位置。零控制输入 q=[0°,0°,0°,0°,0°] 时，末端坐标为 (-28.22, 0, 483.83) mm。该结果与参数链相符：主体沿 z 方向展开，腕部产生 28.22 mm 的 x 向偏移，末端偏距继续沿末端轴方向延伸。",
    )
    add_figure(doc, "图0_实验三实测DH骨架模型.png", "图 1  实验三实测 D-H 参数建立的三维骨架与侧视图", width=6.65)

    doc.add_heading("2.4 实验结果分析", level=2)
    add_body(
        doc,
        "图 1 中各坐标原点按 D-H 链顺序连续连接，底座高度、两段主要臂长、腕部偏移和末端工具长度均在骨架中得到体现。零输入姿态的总高度达到 483.83 mm，主体近似竖直，说明 90° 零位补偿实现了从数学参考方向到工程零位的转换。骨架没有出现关节断链或长度单位突变，表明五组实测参数可以被后续轨迹实验直接复用。",
    )
    add_body(
        doc,
        "误差来源主要包括转轴中心定位误差、卡尺读数误差、舵机外壳与实际轴线之间的装配偏差、夹爪工具中心选取误差以及结构受力后的挠曲。仿真只能验证参数链在数学上的自洽性；若用于实机控制，还需要通过已知标定点测量末端坐标并修正零位与长度参数。",
    )

    doc.add_heading("3 实验四 机械臂末端绘制正方形轨迹仿真", level=1)
    doc.add_heading("3.1 路径规划", level=2)
    add_body(
        doc,
        "实验四复用实验三的 dh_params_servo。指导书给定初始关节角 q₀=[0°,-35°,-100°,20°,0°]，正运动学得到正方形起点 P₀=(297.731899,0,78.102945) mm。路径保持 z=78.102945 mm 不变，依次沿 +y、+x、-y、-x 运动 35 mm，最终回到起点。每条边包含 70 个采样点，四条边共 280 点；相邻边在公共顶点处各保留一个重复点，以体现每条边独立平滑启停。",
    )
    add_table(
        doc,
        ["顶点", "x mm", "y mm", "z mm", "运动方向"],
        [
            ["P0", "297.731899", "0.000000", "78.102945", "+y"],
            ["P1", "297.731899", "35.000000", "78.102945", "+x"],
            ["P2", "332.731899", "35.000000", "78.102945", "-y"],
            ["P3", "332.731899", "0.000000", "78.102945", "-x"],
            ["P4", "297.731899", "0.000000", "78.102945", "闭合"],
        ],
        [0.75, 1.2, 1.1, 1.15, 1.15],
        numeric_cols=[1, 2, 3],
    )

    doc.add_heading("3.2 五次时间缩放", level=2)
    add_body(
        doc,
        "每条边使用归一化时间 t∈[0,1] 和五次多项式生成插值进度。位置、速度和加速度边界条件分别为 s(0)=0、s(1)=1，s′(0)=s′(1)=0，s″(0)=s″(1)=0。由此得到：",
    )
    add_equation(doc, "s(t) = 10t³ − 15t⁴ + 6t⁵")
    add_body(
        doc,
        "目标点由 p(t)=p_start+s(t)(p_end-p_start) 计算。图 2 显示 s(t) 单调从 0 增长到 1，速度和加速度在两个端点均回到零，因此各边在顶点处能够平滑停止并重新启动。",
    )
    add_figure(doc, "图5_五次时间缩放.png", "图 2  五次时间缩放及其速度与加速度", width=5.9)

    doc.add_heading("3.3 带边界数值逆运动学", level=2)
    add_body(
        doc,
        "本实验只约束末端位置，不约束姿态。对第 k 个目标点，最小二乘残差由三维位置误差和关节连续性项组成。连续性权重 λ=0.01，上一轨迹点的逆解作为下一点的初值：",
    )
    add_equation(doc, "r(qₖ) = [ p(qₖ) − p*ₖ ;  λ(qₖ − qₖ₋₁) ]")
    add_body(
        doc,
        "求解器采用 scipy.optimize.least_squares 的 trust-region reflective 方法，并在每次求解后用正运动学重新计算末端位置。若误差超过 1 mm 或求解器失败，程序立即报告目标坐标并停止。",
    )
    add_table(
        doc,
        ["关节", "下限 deg", "上限 deg", "仿真实际范围 deg"],
        [
            ["J1", "-90", "90", "0.000000 ～ 6.704658"],
            ["J2", "-90", "90", "-51.138975 ～ -30.788181"],
            ["J3", "-90", "90", "-90.000000 ～ -69.257565"],
            ["J4", "-90", "20", "-2.368107 ～ 5.573469"],
            ["J5", "-90", "90", "0.000000 ～ 0.000000"],
        ],
        [0.8, 1.1, 1.1, 2.25],
        numeric_cols=[1, 2],
    )
    add_body(
        doc,
        "指导书给定的初始 q3=-100° 低于逆解下限 -90°。程序仍使用原始初始姿态计算几何起点，但将首个最小二乘初值截取到 -90°，随后依靠五自由度模型的位置冗余寻找限位内的等效位置解。第一个目标点最终误差仍小于 10⁻⁶ mm，因此既保留了指导书定义的正方形起点，也保证了所有输出关节角满足限位。",
    )

    doc.add_heading("3.4 仿真流程", level=2)
    for item in [
        "导入实验三参数并建立带 qlim 的 DHRobot 模型。",
        "由 q₀ 正运动学计算末端起点，生成四条五次时间缩放边段。",
        "从限位内初值开始，对 280 个目标点逐点执行带边界最小二乘求解。",
        "逐点检查位置误差，保存关节轨迹、实际末端轨迹和误差序列。",
        "绘制目标路径、实际路径、机械臂关键构型、关节角和误差曲线。",
    ]:
        add_number(doc, item)

    doc.add_heading("3.5 轨迹结果", level=2)
    add_figure(doc, "图1_正方形目标与实际轨迹.png", "图 3  正方形目标轨迹与实际末端轨迹俯视对比", width=5.65)
    add_body(
        doc,
        "图 3 中红色实际轨迹与灰色目标轨迹在显示分辨率下完全重合，四条边长度和运动方向均符合规划，绿色点所示起点与终点重合。路径所有 z 坐标保持为 78.102945 mm，说明正方形位于固定高度的水平面内。",
    )
    add_figure(doc, "图2_机械臂三维轨迹与关键构型.png", "图 4  机械臂沿正方形轨迹运动的关键构型", width=6.2)
    add_body(
        doc,
        "图 4 叠加显示起点、四条边终点附近的机械臂构型以及末端轨迹。随着末端在水平面内移动，底座偏航角和肩肘关节协同变化，末端坐标系方向允许变化，符合“仅约束位置”的实验条件。各构型保持在同一连续运动分支，没有出现突然翻转。",
    )

    doc.add_heading("3.6 精度与连续性分析", level=2)
    add_table(
        doc,
        ["评价指标", "实际结果", "实验要求", "判定"],
        [
            ["有效轨迹点", "280 点", "4×70 点", "满足"],
            ["最大位置误差", "7.9565×10⁻⁷ mm", "≤1 mm", "满足"],
            ["平均位置误差", "8.5152×10⁻⁹ mm", "未规定", "记录"],
            ["实际闭合误差", "7.9564×10⁻⁷ mm", "轨迹闭合", "满足"],
            ["最大相邻关节变化", "0.536826°（J3）", "无明显跳变", "满足"],
            ["J4 最大角度", "5.573469°", "≤20°", "满足"],
        ],
        [1.55, 1.85, 1.55, 0.85],
    )
    add_figure(doc, "图3_逆运动学位置误差.png", "图 5  逐点逆运动学位置误差与 1 mm 阈值", width=5.9)
    add_body(
        doc,
        "图 5 使用对数纵轴显示微小误差。最大值出现在首个目标点，主要原因是首点初值由 -100° 截取到 -90° 后需要寻找限位内的等效构型；其余点以上一解为初值，误差进一步降低。全程误差相对 1 mm 阈值保留约六个数量级的裕量。",
    )
    add_figure(doc, "图4_关节角轨迹.png", "图 6  五个关节角随轨迹采样点的变化", width=6.05)
    add_body(
        doc,
        "图 6 中 J1 至 J4 曲线连续，在四条边的连接处斜率平缓；J5 保持为 0°，说明在位置约束和连续性正则项下无需额外转动末端轴。J3 在部分轨迹点接近 -90° 下限但未越界，J4 的最大值仅为 5.573469°，显著低于 20° 上限。最大相邻关节变化为 0.536826°，没有出现明显关节跳变。",
    )

    doc.add_heading("4 综合讨论", level=1)
    doc.add_heading("4.1 实验三与实验四的参数继承关系", level=2)
    add_body(
        doc,
        "实验四没有重新定义机械臂尺寸，而是直接导入实验三的 dh_params_servo。这样，轨迹可达性、关节运动范围和末端位置全部建立在同一套实测几何模型上。若实验三中的任一长度、扭角或 offset 发生改变，实验四的初始末端位置、逆解和可达工作空间也会随之变化，因此 D-H 标定质量是轨迹仿真的基础。",
    )
    doc.add_heading("4.2 数值结果的含义与局限", level=2)
    add_body(
        doc,
        "本次亚微米级数值误差表示优化器对数学模型目标点的拟合精度，并不代表实体机械臂能够达到同等级定位精度。实机误差还会受到舵机分辨率、回差、连杆挠曲、装配间隙、末端负载和开环控制的影响。若将轨迹下发到实机，应按舵机控制周期对轨迹进行时间标定，并通过实测轨迹重新评价毫米级误差。",
    )
    add_body(
        doc,
        "连续性项抑制了数值逆解在冗余关节之间任意切换，但 λ=0.01 仍属于经验参数。权重过小可能导致构型跳变，权重过大则会牺牲位置精度或造成求解困难。后续可以比较不同 λ 下的最大关节步长、求解次数和位置误差，从而选择更适合实机的折中值。",
    )
    doc.add_heading("4.3 可改进内容", level=2)
    for item in [
        "在实验三中增加多个已知姿态的末端坐标测量，用最小二乘标定连杆长度和零位偏置。",
        "统一规定 D-H 字典的角度单位，并在模型构造函数入口处显式转换和检查。",
        "为轨迹增加关节速度、加速度和舵机速度上限，使采样点能够映射为实际运行时间。",
        "使用编码器或视觉系统记录实际末端路径，将模型误差与机械误差分离。",
        "对接近关节限位或奇异位形的轨迹增加条件数与可操作度检查。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("5 结论", level=1)
    add_body(
        doc,
        "实验三完成了五自由度舵机机械臂的标准 D-H 参数整理和虚拟骨架搭建。零输入姿态的末端位置为 (-28.22, 0, 483.83) mm，模型连续体现了底座高度、主要臂长、腕部偏移和末端工具长度。实验四在同一模型上生成了边长 35 mm、每边 70 点的水平正方形，并通过五次时间缩放、带边界最小二乘逆运动学和连续性约束获得 280 个有效关节配置。",
    )
    add_body(
        doc,
        "仿真最大位置误差为 7.9565×10⁻⁷ mm，闭合误差为 7.9564×10⁻⁷ mm，最大相邻关节变化为 0.536826°，J4 最大角度为 5.573469°。目标轨迹与实际轨迹一致，所有关节均满足限位，验证了实测 D-H 模型、笛卡尔路径规划和连续数值逆运动学方案的正确性。",
    )

    doc.add_heading("参考资料", level=1)
    refs = [
        "第五周实验指导书：实验三 舵机机械臂实机 D-H 参数测量与虚拟搭建。",
        "第五周实验指导书：实验四 机械臂末端绘制正方形轨迹仿真。",
        "Peter Corke, Robotics, Vision and Control, Springer。",
        "Robotics Toolbox for Python 文档及 SciPy least_squares 接口文档。",
    ]
    for index, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        r = p.add_run(f"{index}. {ref}")
        set_run_font(r)

    # 文档属性
    doc.core_properties.title = "五自由度舵机机械臂运动学建模与正方形轨迹仿真实验报告"
    doc.core_properties.subject = "实验三与实验四"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "D-H, Robotics Toolbox, 轨迹规划, 逆运动学"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
