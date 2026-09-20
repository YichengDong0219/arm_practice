from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\Code\arm_pracctice")
TEMPLATE = Path(r"E:\Desktop\机械臂实践\26秋季初级实践课实验报告模板.docx")
OUTPUT = ROOT / "机械臂平面运动与圆周轨迹实验报告.docx"
ASSET_DIR = ROOT / ".tmp_report" / "assets"


def set_run_font(run, *, size=12, bold=False, east_asia="宋体", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.25, first_line=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Pt(24)


def replace_text_preserve_paragraph(paragraph, text, *, size=12, bold=False, east_asia="宋体"):
    # 模板日期段中含有域/书签节点，单纯删除 runs 会残留“×××”占位符。
    p_element = paragraph._element
    for child in list(p_element):
        if child.tag != qn("w:pPr"):
            p_element.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, east_asia=east_asia)


def add_heading(doc, text, *, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        set_paragraph_spacing(p, before=10, after=5, line=1.0)
        size = 12
    else:
        set_paragraph_spacing(p, before=7, after=3, line=1.0)
        size = 12
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True)
    return p


def add_body(doc, text, *, indent=True, after=5):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=after, line=1.25, first_line=indent)
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    set_run_font(r, size=12)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        set_paragraph_spacing(p, before=0, after=2, line=1.15)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        r = p.add_run(f"•  {item}")
        set_run_font(r, size=11)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    set_paragraph_spacing(p, before=2, after=6, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_picture(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=3, after=0, line=1.0)
    p.add_run().add_picture(str(path), width=width)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths_cm=None, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "E7E6E6")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        if widths_cm:
            cell.width = Cm(widths_cm[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, line=1.0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, bold=True)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths_cm:
                cell.width = Cm(widths_cm[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, line=1.08)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    return table


def remove_from_paragraph_to_end(doc, start_paragraph):
    body = doc._element.body
    start = start_paragraph._element
    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if deleting and child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体")


def build_report():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"未找到实验报告模板：{TEMPLATE}")
    doc = Document(str(TEMPLATE))
    configure_styles(doc)

    # 保留模板的 A4 页面设置、页边距和基本信息表，只替换标题与日期。
    title_p = doc.paragraphs[0]
    title_p.style = doc.styles["Title"]
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title_p, before=0, after=12, line=1.0)
    replace_text_preserve_paragraph(
        title_p,
        "机器人实践课机械臂平面运动与圆周轨迹实验报告",
        size=16,
        bold=True,
        east_asia="楷体",
    )

    for p in doc.paragraphs:
        if p.text.strip().startswith("实验日期"):
            replace_text_preserve_paragraph(p, "实验日期：2026年9月12日", size=12)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_spacing(p, before=0, after=5, line=1.0)
            break

    first_experiment = next(p for p in doc.paragraphs if p.text.strip().startswith("实验1"))
    remove_from_paragraph_to_end(doc, first_experiment)

    add_heading(doc, "项目概述", level=1)
    add_body(
        doc,
        "本工程围绕6自由度总线舵机机械臂在锁定Joint 1后的二维平面运动展开。工程从单关节平滑控制起步，完成Joint 2、3、4的顺序与同步驱动，随后把关节空间插值扩展为笛卡尔圆轨迹规划。最终版本以Joint 2为原点，采用L1=120 mm、L2=120 mm、L3=140 mm的平面3R模型，使Joint 4上方140 mm处的末端点围绕零位竖直线上的圆心运动，并通过反向构型避开Joint 4的机械限位。当前已完成协议、可达性、连续性和理论圆度验证，实机圆度与重复定位精度仍待测量。",
    )
    add_table(
        doc,
        ["文件", "主要作用", "完成状态"],
        [
            ["task6.py", "Joint 4单轴正反向平滑运动，验证基础串口控制", "已完成代码与协议检查"],
            ["task7.py", "Joint 2、3、4依次运动及三轴同步线性插值", "已完成代码与轨迹逻辑检查"],
            ["task8.py", "二连杆模型下的Joint 4轴心圆周运动原型", "保留为原始方案"],
            ["task8_reverse.py", "三连杆末端圆周运动、半径/圈数接口及反向IK", "当前推荐实机版本"],
        ],
        widths_cm=[3.1, 9.0, 4.0],
    )

    add_heading(doc, "实验1：关节空间插值与多舵机协同驱动", level=1)
    add_heading(doc, "1.1 过程记录", level=2)
    add_body(
        doc,
        "首先处理Python运行环境：Windows终端中的python3实际指向MSYS2解释器，未安装pyserial，因而出现“ModuleNotFoundError: No module named 'serial'”。改为激活ra_class环境后使用python运行，并通过python -m pip安装pyserial与numpy。task6沿用STM32固定16字节帧，控制Joint 4按0°→45°→−30°→0°运动；每段1.5 s、30个插值点。task7继续采用相同协议，依次将Joint 2、3、4置于30°、−45°、15°，再三轴同步移动至45°、−30°、−15°并复位。",
    )
    add_table(
        doc,
        ["项目", "设置或计算结果"],
        [
            ["通信", "COM5，115200 bps，timeout=1 s；打开前DTR=False、RTS=False"],
            ["协议", "16字节：0xAA + 6×有符号大端角度 + 模式0x01 + XOR + 0xBB"],
            ["角度编码", "弧度转换为0.1°单位的有符号16位整数"],
            ["task6轨迹", "3段×30帧，共90个控制帧"],
            ["task7轨迹", "3段单轴×30帧 + 2段联动×40帧，共170个控制帧"],
        ],
        widths_cm=[3.7, 12.4],
    )

    add_heading(doc, "1.2 结果与讨论", level=2)
    add_body(
        doc,
        "task6和task7的运动本质都是关节空间线性插值：每个周期按统一时间比例计算目标角并下发，因此能够限制相邻帧角度突变、降低电流冲击并保证多轴同时启停。task6仅改变pose[3]；task7通过复制姿态数组并分别修改pose[1]、pose[2]、pose[3]实现单轴保持与多轴联动。该方法适合验证关节关系，但末端路径由正运动学自然产生，并不是预先指定的直线或圆。代码没有关节位置反馈，启动姿态只能假定为全零位；若真实姿态不一致，首次插值仍可能产生突跳。",
    )
    add_heading(doc, "1.3 团队分工情况", level=2)
    add_body(doc, "报告人：________________；同组队员1：________________；同组队员2：________________。", indent=False)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "实验2：平面三连杆末端圆周轨迹与冗余逆运动学", level=1)
    add_heading(doc, "2.1 方案设计与迭代过程", level=2)
    add_body(
        doc,
        "为让末端在平面内画圆，先在笛卡尔空间离散圆周坐标，再逐点求逆运动学并按50 Hz发送。原始task8把Joint 4轴心视为二连杆末端；需求进一步明确后，将Joint 4上方140 mm处的实际末端纳入模型，形成L1=120 mm、L2=120 mm、L3=140 mm的3R机构。由于位置任务只有两个约束而有三个可动关节，利用一个冗余自由度搜索可行末端方向：每个圆周点遍历方向偏置，求负肘2R解，再计算q4，并以限位、连续性和最大关节角作为筛选条件。",
    )
    add_picture(doc, ASSET_DIR / "equations.png", Inches(6.05), "图1  关节插值、三连杆正运动学与冗余IK降维关系")
    add_table(
        doc,
        ["迭代阶段", "发现的问题", "处理方法与结果"],
        [
            ["二连杆原型", "只控制Joint 4轴心，未覆盖其上方实际末端", "建立2R解析IK，验证圆周点逐点求解流程"],
            ["参数接口", "半径和连续圈数写死，不便实机调试", "加入--radius/-r与--cycles/-n命令行接口"],
            ["三连杆修正", "140 mm末端段被误解为新增电机", "确认仍只动Joint 2、3、4，改为3R位置IK"],
            ["零位修正", "早期模型把全零方向设为水平，圆心位置不符", "加入90°零位偏置，使全零连杆位于x=0竖直线；圆心取(0,240 mm)"],
            ["关节4限位", "正向解会进入物理受限方向；简单取反会破坏运动学", "使用负肘支路并搜索末端方向，使q4全程为负且保持轨迹闭合"],
            ["姿态约束", "固定末端姿态压缩可达空间", "取消固定姿态，将q4作为冗余变量参与求解，扩大位置可行域"],
        ],
        widths_cm=[3.1, 5.2, 7.8],
        font_size=9.5,
    )

    add_heading(doc, "2.2 最终模型、验证结果与讨论", level=2)
    add_picture(doc, ASSET_DIR / "endpoint_circle.png", Inches(5.65), "图2  三连杆零位竖直线、圆心位置与末端圆周轨迹")
    add_body(
        doc,
        "当前task8_reverse.py默认圆心为(0,240 mm)，位于三个关节全零位组成的竖直线上；默认半径18 mm、8 s/圈、50 Hz，每圈400个控制点。程序先离线生成完整轨迹，再统一检查有限数值、可达性、±90°软件限位、非受限转向、相邻帧连续性、闭环一致性和协议校验；任一检查失败时在打开串口前拒绝运动。正常流程为2 s到达圆起点、保持0.5 s、运行n圈、保持0.5 s并用2 s复位；中断或串口异常时停止发送并关闭串口，不强制复位。",
    )
    add_table(
        doc,
        ["验证项", "默认18 mm半径的结果"],
        [
            ["连杆与圆心", "L1/L2/L3=120/120/140 mm；圆心(0,240 mm)"],
            ["选定方向偏置", "−61°"],
            ["Joint 2范围", "51.706° ～ 74.384°"],
            ["Joint 3范围", "−69.792° ～ −39.128°"],
            ["Joint 4范围", "−74.219° ～ −64.569°，全程保持反向转动区间"],
            ["最大相邻帧变化", "约0.239°，未出现构型跳变"],
            ["正运动学圆度误差", "最大约1.42×10⁻¹³ mm（浮点计算误差量级）"],
            ["半径可用性预演", "18、30、40 mm通过；50 mm因限位/可达性检查被拒绝"],
            ["两圈发送预演", "总计880帧：40帧进场 + 800帧圆周 + 40帧复位"],
        ],
        widths_cm=[4.8, 11.3],
    )
    add_picture(doc, ASSET_DIR / "joint_angles.png", Inches(6.05), "图3  默认18 mm圆周轨迹下Joint 2、3、4的角度变化")
    add_body(
        doc,
        "上述结果来自无硬件数学预演与伪串口全流程测试，可证明IK方程、轨迹连续性和数据帧逻辑自洽，但不能等同于实机轨迹精度。实际圆度还会受到舵机零位偏差、齿隙、结构挠曲、安装尺寸误差、负载和控制周期抖动影响。后续应从较小半径和单圈开始实机验收，记录末端轨迹与关节温升，再逐步增加半径和圈数；若实测圆心偏移，应优先标定三段有效长度、零位偏置及舵机正方向。",
    )
    add_heading(doc, "2.3 使用接口与团队分工情况", level=2)
    add_body(doc, "运行示例：python task8_reverse.py --radius 30 --cycles 3。半径单位为mm，圈数必须为正整数；默认值分别为18和1。", indent=False)
    add_body(doc, "报告人：________________；同组队员1：________________；同组队员2：________________。", indent=False)

    add_heading(doc, "结论", level=1)
    add_body(
        doc,
        "工程已形成从串口协议、关节插值到笛卡尔轨迹和冗余IK的完整技术路线。task8_reverse.py是当前推荐方案，已通过软件验证并提供半径、连续圈数接口；后续应完成实机标定与轨迹测量。",
    )

    # 全文孤行控制；标题与其后内容尽量保持同页。
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.core_properties.title = "机械臂平面运动与圆周轨迹实验报告"
    doc.core_properties.subject = "task6、task7、task8及task8_reverse工程总结"
    doc.core_properties.author = ""
    doc.core_properties.comments = "基于课程实验报告模板生成；实机数据栏未虚构。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
