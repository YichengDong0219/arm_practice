from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


ROOT = Path(r"E:\Code\arm_pracctice")
GUIDE = ROOT / "第四周实验指导书(2).docx"
OUT = ROOT / ".tmp_week4"
MEDIA = OUT / "media"
OUT.mkdir(exist_ok=True)
MEDIA.mkdir(exist_ok=True)

doc = Document(str(GUIDE))

paragraphs = []
for index, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else "",
                "text": text,
            }
        )

tables = []
for table_index, table in enumerate(doc.tables, start=1):
    rows = []
    for row in table.rows:
        rows.append([cell.text.replace("\n", " | ").strip() for cell in row.cells])
    tables.append({"index": table_index, "rows": rows})

media_files = []
with zipfile.ZipFile(GUIDE) as archive:
    for name in archive.namelist():
        if name.startswith("word/media/") and not name.endswith("/"):
            target = MEDIA / Path(name).name
            target.write_bytes(archive.read(name))
            media_files.append(
                {
                    "archive_name": name,
                    "path": str(target),
                    "size": target.stat().st_size,
                }
            )

result = {
    "guide": str(GUIDE),
    "paragraphs": paragraphs,
    "tables": tables,
    "inline_shapes": len(doc.inline_shapes),
    "media": media_files,
    "sections": [
        {
            "width_in": round(section.page_width.inches, 3),
            "height_in": round(section.page_height.inches, 3),
            "top_margin_in": round(section.top_margin.inches, 3),
            "bottom_margin_in": round(section.bottom_margin.inches, 3),
            "left_margin_in": round(section.left_margin.inches, 3),
            "right_margin_in": round(section.right_margin.inches, 3),
        }
        for section in doc.sections
    ],
}

# 按文档真实顺序建立“实验章节 -> 示例代码表”映射。
ordered_blocks = []
current_experiment = ""
table_number = 0
for child in doc.element.body.iterchildren():
    if isinstance(child, CT_P):
        paragraph = Paragraph(child, doc)
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        image_refs = []
        for blip in paragraph._p.xpath(".//a:blip"):
            rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if rel_id and rel_id in doc.part.rels:
                image_refs.append(Path(doc.part.rels[rel_id].target_ref).name)
        if style == "Heading 2" and text.startswith("1."):
            current_experiment = text
        if text or image_refs:
            ordered_blocks.append(
                {
                    "type": "paragraph",
                    "style": style,
                    "text": text,
                    "experiment": current_experiment,
                    "images": image_refs,
                }
            )
    elif isinstance(child, CT_Tbl):
        table_number += 1
        table = Table(child, doc)
        raw_cells = [[cell.text for cell in row.cells] for row in table.rows]
        ordered_blocks.append(
            {
                "type": "table",
                "table_index": table_number,
                "experiment": current_experiment,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "cell_text": raw_cells,
            }
        )

result["ordered_blocks"] = ordered_blocks

example_dir = OUT / "guide_examples_raw"
example_dir.mkdir(exist_ok=True)
for block in ordered_blocks:
    if block["type"] == "table" and 1 <= block["table_index"] <= 10:
        code = block["cell_text"][0][0]
        (example_dir / f"experiment_{block['table_index']:02d}.py").write_text(
            code.rstrip() + "\n", encoding="utf-8"
        )

(OUT / "guide_structure.json").write_text(
    json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
)

with (OUT / "guide_content.txt").open("w", encoding="utf-8") as handle:
    for item in paragraphs:
        handle.write(f"[P{item['index']}][{item['style']}] {item['text']}\n")
    for table in tables:
        handle.write(f"\n[TABLE {table['index']}]\n")
        for row in table["rows"]:
            handle.write(" || ".join(row) + "\n")

print(
    json.dumps(
        {
            "paragraphs": len(paragraphs),
            "tables": len(tables),
            "inline_shapes": len(doc.inline_shapes),
            "media": len(media_files),
        },
        ensure_ascii=False,
    )
)
